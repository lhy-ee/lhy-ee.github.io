from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUTPUT = r"A:\temp\profilepage\_test\CV_image_reference.docx"
FONT = "Times New Roman"

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.27)
section.page_height = Inches(11.69)
section.top_margin = Inches(0.46)
section.bottom_margin = Inches(0.46)
section.left_margin = Inches(0.58)
section.right_margin = Inches(0.58)

def font(style, size, bold=False, italic=False):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    style.font.color.rgb = RGBColor(0, 0, 0)
    style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

normal = doc.styles["Normal"]
font(normal, 8.35)
normal.paragraph_format.space_after = Pt(0.8)
normal.paragraph_format.line_spacing = 0.98

title = doc.styles["Title"]
font(title, 18, True)
title.paragraph_format.space_after = Pt(1)
title.paragraph_format.keep_with_next = True

subtitle = doc.styles.add_style("Reference Subtitle", WD_STYLE_TYPE.PARAGRAPH)
font(subtitle, 9.1, True)
subtitle.paragraph_format.space_after = Pt(6)
subtitle.paragraph_format.keep_with_next = True

section_style = doc.styles.add_style("Reference Section", WD_STYLE_TYPE.PARAGRAPH)
font(section_style, 10.6, True)
section_style.paragraph_format.space_before = Pt(5.8)
section_style.paragraph_format.space_after = Pt(1.4)
section_style.paragraph_format.keep_with_next = True

entry = doc.styles.add_style("Reference Entry", WD_STYLE_TYPE.PARAGRAPH)
font(entry, 8.7, True)
entry.paragraph_format.space_after = Pt(0.4)
entry.paragraph_format.keep_with_next = True

meta = doc.styles.add_style("Reference Meta", WD_STYLE_TYPE.PARAGRAPH)
font(meta, 8.1, False, True)
meta.paragraph_format.space_after = Pt(1)
meta.paragraph_format.keep_with_next = True

bullet = doc.styles.add_style("Reference Bullet", WD_STYLE_TYPE.PARAGRAPH)
font(bullet, 8.15)
bullet.paragraph_format.left_indent = Inches(0.22)
bullet.paragraph_format.first_line_indent = Inches(-0.15)
bullet.paragraph_format.space_after = Pt(0.8)
bullet.paragraph_format.line_spacing = 0.96

def keep(p, next_paragraph=False):
    ppr = p._p.get_or_add_pPr()
    ppr.append(OxmlElement("w:keepLines"))
    if next_paragraph:
        ppr.append(OxmlElement("w:keepNext"))

def border_bottom(p):
    ppr = p._p.get_or_add_pPr()
    borders = ppr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        ppr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    borders.append(bottom)

def add_section(text):
    p = doc.add_paragraph(text, style="Reference Section")
    border_bottom(p)
    keep(p, True)

def add_three_column(left, center="", right=""):
    p = doc.add_paragraph(style="Reference Entry")
    p.paragraph_format.tab_stops.add_tab_stop(Inches(3.55), WD_TAB_ALIGNMENT.CENTER)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(7.05), WD_TAB_ALIGNMENT.RIGHT)
    p.add_run(left).bold = True
    if center:
        r = p.add_run("\t" + center)
        r.bold = False
        r.italic = True
    if right:
        p.add_run("\t" + right).italic = True
    keep(p, True)
    return p

def add_bullet(label, text):
    p = doc.add_paragraph(style="Reference Bullet")
    p.add_run("•  ")
    if label:
        p.add_run(label).bold = True
        p.add_run(": " + text)
    else:
        p.add_run(text)
    keep(p)

name = doc.add_paragraph("LI HONG YI", style="Title")
name.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub = doc.add_paragraph("Technical Artist  |  Creative Technologist  |  3D Artist", style="Reference Subtitle")
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

add_section("EDUCATION")
add_three_column("Sichuan Fine Arts Institute", "Bachelor of Fine Arts in Film and Television Production", "2018-2022")
p = doc.add_paragraph()
p.add_run("Chongqing, China").italic = True
p.add_run("  |  ")
p.add_run("GPA: ").bold = True
p.add_run("3.22")
keep(p, True)
p = doc.add_paragraph()
p.add_run("Core Courses: ").bold = True
p.add_run("Film Theory, Audiovisual Language, Film and Television Production, Digital Media Production, 3D Visualization, Virtual Visual Effects")
keep(p)

add_section("PROFESSIONAL EXPERIENCE")
add_three_column("Tencent Games - LIGHTSPEED STUDIOS", "Applied R&D (Technical Artist)", "Aug 2023-Mar 2026")
p = doc.add_paragraph("Omni Research  |  Shenzhen Tengyu Interactive Technology Co., Ltd. (a Tencent subsidiary)", style="Reference Meta")
keep(p, True)
add_bullet("Batch-processing pipelines", "Contributed to the design and continuous iteration of pipelines for 3D data in AI auto-rigging projects.")
add_bullet("Data production and quality control", "Managed task planning, team training, quality review, issue tracking, and final delivery.")
add_bullet("AI auto-rigging evaluation", "Analyzed and categorized recurring issues across 3D models of varying complexity and contributed to post-processing features and plugin interactions.")
add_bullet("AI tool evaluation and documentation", "Evaluated capabilities, performance, and interaction workflows; created and maintained data-production guidelines, workflow documentation, and tool documentation.")
add_bullet("Visual-content support", "Provided 3D modeling, animation processing, rendering, and related support for multiple AI research projects.")

p = doc.add_paragraph("Selected Contributions", style="Reference Meta")
p.runs[0].bold = True
add_bullet("Peacekeeper Elite (和平精英)", "Developed and continuously iterated batch-processing pipelines for an AI auto-rigging project; evaluated internal AI tools and contributed to post-processing features, plugin interactions, and integration into practical art-production workflows.")
add_bullet("Aoxing Heatwave (奥星热浪)", "Managed motion-capture data production for an AI-generated dance project, including requirement coordination, annotation-team management, quality control, data delivery, and cross-team collaboration.")

add_three_column("Chongqing Yujitu Technology Co., Ltd.", "3D Artist", "Oct 2022-Apr 2023")
p = doc.add_paragraph("重庆禹迹图科技有限公司", style="Reference Meta")
keep(p, True)
add_bullet("3D asset production", "Created production-ready 3D models for commercial projects, following project specifications and asset requirements.")

add_section("RESEARCH CONTRIBUTIONS")
add_three_column("Light-SQ", "SIGGRAPH Asia 2025  |  Co-author", "")
p = doc.add_paragraph("Structure-Aware Shape Abstraction with Superquadrics for Generated Meshes", style="Reference Meta")
keep(p, True)
add_bullet("Technical contribution", "Contributed to the development of a batch-rendering script.")
add_bullet("Visual contribution", "Designed and produced the project cover and poster.")

add_three_column("TailorRig", "Research Contribution", "")
p = doc.add_paragraph("Controllable and Structured Auxiliary Skeleton Generation for Garment-Aware Rigging", style="Reference Meta")
keep(p, True)
add_bullet("Dataset development", "Contributed continuously to dataset development and data production.")
add_bullet("Visual contribution", "Designed and produced the project cover, poster, and related visual materials.")

add_section("TECHNICAL SKILLS")
for label, text in [
    ("Skills", "3D Modeling, Procedural Materials, Real-Time Rendering, Shader Development, Data Processing, Pipeline Development, Workflow Automation"),
    ("Software", "Unreal Engine, Blender, Substance 3D Designer, Substance 3D Painter, Marvelous Designer, ZBrush, MotionBuilder"),
    ("Programming Languages", "Python, HLSL"),
]:
    p = doc.add_paragraph()
    p.add_run(label + ": ").bold = True
    p.add_run(text)
    keep(p)

add_section("RESEARCH INTERESTS")
p = doc.add_paragraph("Human-AI Collaboration, Creative Tools, Computational Media, Authoring Systems, Computer Graphics, Procedural Materials, Real-Time Rendering")
keep(p)

for p in doc.paragraphs:
    p.paragraph_format.widow_control = True

doc.core_properties.title = "Li Hong Yi Curriculum Vitae"
doc.core_properties.subject = "Technical Artist Curriculum Vitae"
doc.save(OUTPUT)
print(OUTPUT)

