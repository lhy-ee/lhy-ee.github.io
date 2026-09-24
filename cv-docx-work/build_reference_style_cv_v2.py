from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUTPUT = r"A:\temp\profilepage\_test\CV_image_reference_v2.docx"
FONT = "Times New Roman"

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.27), Inches(11.69)
section.top_margin, section.bottom_margin = Inches(0.44), Inches(0.44)
section.left_margin, section.right_margin = Inches(0.58), Inches(0.58)

def set_run(run, size=None, bold=None, italic=None):
    run.font.name = FONT
    run.font.color.rgb = RGBColor(0, 0, 0)
    if size is not None: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    fonts.set(qn("w:ascii"), FONT)
    fonts.set(qn("w:hAnsi"), FONT)
    fonts.set(qn("w:eastAsia"), "SimSun")

def set_style(style, size, bold=False, italic=False):
    style.font.name, style.font.size = FONT, Pt(size)
    style.font.bold, style.font.italic = bold, italic
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), FONT); fonts.set(qn("w:hAnsi"), FONT); fonts.set(qn("w:eastAsia"), "SimSun")

normal = doc.styles["Normal"]
set_style(normal, 8.0)
normal.paragraph_format.space_after = Pt(0.35)
normal.paragraph_format.line_spacing = 0.94

section_style = doc.styles.add_style("CV Section", WD_STYLE_TYPE.PARAGRAPH)
set_style(section_style, 9.7, True)
section_style.paragraph_format.space_before = Pt(4.4)
section_style.paragraph_format.space_after = Pt(0.7)
section_style.paragraph_format.keep_with_next = True

meta_style = doc.styles.add_style("CV Meta", WD_STYLE_TYPE.PARAGRAPH)
set_style(meta_style, 7.75, False, True)
meta_style.paragraph_format.space_after = Pt(0.35)
meta_style.paragraph_format.keep_with_next = True

bullet_style = doc.styles.add_style("CV Bullet", WD_STYLE_TYPE.PARAGRAPH)
set_style(bullet_style, 7.75)
bullet_style.paragraph_format.left_indent = Inches(0.23)
bullet_style.paragraph_format.first_line_indent = Inches(-0.15)
bullet_style.paragraph_format.space_after = Pt(0.45)
bullet_style.paragraph_format.line_spacing = 0.92

def keep(p, next_=False):
    ppr = p._p.get_or_add_pPr()
    if ppr.find(qn("w:keepLines")) is None: ppr.append(OxmlElement("w:keepLines"))
    if next_ and ppr.find(qn("w:keepNext")) is None: ppr.append(OxmlElement("w:keepNext"))

def add_section(text):
    p = doc.add_paragraph(style="CV Section")
    r = p.add_run(text); set_run(r, 9.7, True, False)
    ppr = p._p.get_or_add_pPr(); borders = OxmlElement("w:pBdr"); bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "5"); bottom.set(qn("w:space"), "0"); bottom.set(qn("w:color"), "000000")
    borders.append(bottom); ppr.append(borders); keep(p, True)

def set_cell_margins(cell, top=0, start=0, bottom=0, end=0):
    tcpr = cell._tc.get_or_add_tcPr(); tcMar = tcpr.first_child_found_in("w:tcMar")
    if tcMar is None: tcMar = OxmlElement("w:tcMar"); tcpr.append(tcMar)
    for side, value in (("top",top),("start",start),("bottom",bottom),("end",end)):
        node = OxmlElement("w:"+side); node.set(qn("w:w"), str(value)); node.set(qn("w:type"), "dxa"); tcMar.append(node)

def remove_table_borders(table):
    tblpr = table._tbl.tblPr; borders = OxmlElement("w:tblBorders")
    for edge in ("top","left","bottom","right","insideH","insideV"):
        e = OxmlElement("w:"+edge); e.set(qn("w:val"), "nil"); borders.append(e)
    tblpr.append(borders)

def add_three_column(left, center="", right=""):
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT; table.autofit = False
    widths = (Inches(3.0), Inches(2.72), Inches(1.39))
    remove_table_borders(table)
    for cell, width in zip(table.rows[0].cells, widths):
        cell.width = width; cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_margins(cell)
        p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 0.9; keep(p, True)
    vals = ((left, WD_ALIGN_PARAGRAPH.LEFT, True, False), (center, WD_ALIGN_PARAGRAPH.CENTER, False, True), (right, WD_ALIGN_PARAGRAPH.RIGHT, False, True))
    for cell, (text, align, bold, italic) in zip(table.rows[0].cells, vals):
        p = cell.paragraphs[0]; p.alignment = align
        if text: set_run(p.add_run(text), 8.1, bold, italic)
    return table

def add_bullet(label, text):
    p = doc.add_paragraph(style="CV Bullet")
    set_run(p.add_run("•  "), 7.75)
    if label:
        set_run(p.add_run(label), 7.75, True, False); set_run(p.add_run(": " + text), 7.75, False, False)
    else: set_run(p.add_run(text), 7.75)
    keep(p)

def add_line(parts, style=None, next_=False):
    p = doc.add_paragraph(style=style)
    for text, bold, italic in parts: set_run(p.add_run(text), 7.85, bold, italic)
    keep(p, next_); return p

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(0)
set_run(p.add_run("LI HONG YI"), 16.5, True, False); keep(p, True)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(4.2)
set_run(p.add_run("Technical Artist  |  Creative Technologist  |  3D Artist"), 8.3, True, False); keep(p, True)

add_section("EDUCATION")
add_three_column("Sichuan Fine Arts Institute", "Bachelor of Fine Arts in Film and Television Production", "2018–2022")
add_line([("Chongqing, China", False, True), ("  |  ", False, False), ("GPA: ", True, False), ("3.22", False, False)], next_=True)
add_line([("Core Courses: ", True, False), ("Film Theory, Audiovisual Language, Film and Television Production, Digital Media Production, 3D Visualization, Virtual Visual Effects", False, False)])

add_section("PROFESSIONAL EXPERIENCE")
add_three_column("Tencent Games – LIGHTSPEED STUDIOS", "Applied R&D (Technical Artist)", "Aug 2023–Mar 2026")
add_line([("Omni Research  |  Shenzhen Tengyu Interactive Technology Co., Ltd. (a Tencent subsidiary)", False, True)], "CV Meta", True)
add_bullet("Batch-processing pipelines", "Contributed to the design and continuous iteration of pipelines for 3D data in AI auto-rigging projects.")
add_bullet("Data production and quality control", "Managed task planning, team training, quality review, issue tracking, and final delivery.")
add_bullet("AI auto-rigging evaluation", "Analyzed and categorized recurring issues across 3D models of varying complexity and contributed to post-processing features and plugin interactions.")
add_bullet("AI tool evaluation and documentation", "Evaluated capabilities, performance, and interaction workflows; created and maintained data-production guidelines, workflow documentation, and tool documentation.")
add_bullet("Visual-content support", "Provided 3D modeling, animation processing, rendering, and related support for multiple AI research projects.")
add_line([("Selected Contributions", True, True)], "CV Meta", True)
add_bullet("Peacekeeper Elite (和平精英)", "Developed and continuously iterated batch-processing pipelines for an AI auto-rigging project; evaluated internal AI tools and contributed to post-processing features, plugin interactions, and integration into practical art-production workflows.")
add_bullet("Aoxing Heatwave (奥星热浪)", "Managed motion-capture data production for an AI-generated dance project, including requirement coordination, annotation-team management, quality control, data delivery, and cross-team collaboration.")
add_three_column("Chongqing Yujitu Technology Co., Ltd.", "3D Artist", "Oct 2022–Apr 2023")
add_line([("重庆禹迹图科技有限公司", False, True)], "CV Meta", True)
add_bullet("3D asset production", "Created production-ready 3D models for commercial projects, following project specifications and asset requirements.")

add_section("RESEARCH CONTRIBUTIONS")
add_three_column("Light-SQ", "SIGGRAPH Asia 2025  |  Co-author", "")
add_line([("Structure-Aware Shape Abstraction with Superquadrics for Generated Meshes", False, True)], "CV Meta", True)
add_bullet("Technical contribution", "Contributed to the development of a batch-rendering script.")
add_bullet("Visual contribution", "Designed and produced the project cover and poster.")
add_three_column("TailorRig", "Research Contribution", "")
add_line([("Controllable and Structured Auxiliary Skeleton Generation for Garment-Aware Rigging", False, True)], "CV Meta", True)
add_bullet("Dataset development", "Contributed continuously to dataset development and data production.")
add_bullet("Visual contribution", "Designed and produced the project cover, poster, and related visual materials.")

add_section("TECHNICAL SKILLS")
for label, value in [
    ("Skills", "3D Modeling, Procedural Materials, Real-Time Rendering, Shader Development, Data Processing, Pipeline Development, Workflow Automation"),
    ("Software", "Unreal Engine, Blender, Substance 3D Designer, Substance 3D Painter, Marvelous Designer, ZBrush, MotionBuilder"),
    ("Programming Languages", "Python, HLSL")]:
    add_line([(label + ": ", True, False), (value, False, False)])

add_section("RESEARCH INTERESTS")
add_line([("Human-AI Collaboration, Creative Tools, Computational Media, Authoring Systems, Computer Graphics, Procedural Materials, Real-Time Rendering", False, False)])

for p in doc.paragraphs:
    p.paragraph_format.widow_control = True
doc.core_properties.title = "Li Hong Yi Curriculum Vitae"
doc.core_properties.subject = "Technical Artist Curriculum Vitae"
doc.save(OUTPUT)
print(OUTPUT)
