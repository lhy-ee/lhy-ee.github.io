from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

SOURCE = r"A:\temp\profilepage\profilepage_test\CV.docx"
OUTPUT = r"A:\temp\profilepage\_test\CV_reformatted.docx"

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.27)
section.page_height = Inches(11.69)
section.top_margin = Inches(0.55)
section.bottom_margin = Inches(0.55)
section.left_margin = Inches(0.68)
section.right_margin = Inches(0.68)

styles = doc.styles

def set_font(style, name="Arial", size=Pt(9.4), bold=False, color="000000"):
    style.font.name = name
    style.font.size = size
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    style._element.rPr.rFonts.set(qn("w:ascii"), name)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), name)

normal = styles["Normal"]
set_font(normal)
normal.paragraph_format.space_after = Pt(2.8)
normal.paragraph_format.line_spacing = 1.08

title = styles["Title"]
set_font(title, size=Pt(24), bold=True)
title.paragraph_format.space_after = Pt(2)
title.paragraph_format.keep_with_next = True

for name in ["Heading 1", "Heading 2", "Heading 3"]:
    st = styles[name]
    st.paragraph_format.keep_with_next = True
    st.paragraph_format.keep_together = True

set_font(styles["Heading 1"], size=Pt(12.5), bold=True)
styles["Heading 1"].paragraph_format.space_before = Pt(9)
styles["Heading 1"].paragraph_format.space_after = Pt(4)

set_font(styles["Heading 2"], size=Pt(10.4), bold=True)
styles["Heading 2"].paragraph_format.space_before = Pt(5)
styles["Heading 2"].paragraph_format.space_after = Pt(1.5)

set_font(styles["Heading 3"], size=Pt(9.2), bold=True, color="333333")
styles["Heading 3"].paragraph_format.space_before = Pt(3.5)
styles["Heading 3"].paragraph_format.space_after = Pt(1.5)

if "CV Subtitle" not in styles:
    subtitle = styles.add_style("CV Subtitle", WD_STYLE_TYPE.PARAGRAPH)
else:
    subtitle = styles["CV Subtitle"]
set_font(subtitle, size=Pt(10.5), color="333333")
subtitle.paragraph_format.space_after = Pt(8)
subtitle.paragraph_format.keep_with_next = True

if "CV Meta" not in styles:
    meta = styles.add_style("CV Meta", WD_STYLE_TYPE.PARAGRAPH)
else:
    meta = styles["CV Meta"]
set_font(meta, size=Pt(8.8), color="444444")
meta.paragraph_format.space_after = Pt(2.5)
meta.paragraph_format.keep_with_next = True

if "CV Bullet" not in styles:
    bullet = styles.add_style("CV Bullet", WD_STYLE_TYPE.PARAGRAPH)
else:
    bullet = styles["CV Bullet"]
set_font(bullet, size=Pt(9.15))
bullet.paragraph_format.left_indent = Inches(0.18)
bullet.paragraph_format.first_line_indent = Inches(-0.13)
bullet.paragraph_format.space_after = Pt(2.4)
bullet.paragraph_format.line_spacing = 1.05

def keep(paragraph, next_one=False):
    ppr = paragraph._p.get_or_add_pPr()
    if next_one:
        ppr.append(OxmlElement("w:keepNext"))
    ppr.append(OxmlElement("w:keepLines"))

def add_section(title_text):
    p = doc.add_paragraph(title_text, style="Heading 1")
    return p

def add_role(left, right=None, organization=None, meta_text=None):
    p = doc.add_paragraph(style="Heading 2")
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.7), WD_TAB_ALIGNMENT.RIGHT)
    p.add_run(left).bold = True
    if organization:
        p.add_run(" | " + organization).bold = False
    if right:
        p.add_run("\t" + right).bold = True
    keep(p, True)
    if meta_text:
        m = doc.add_paragraph(meta_text, style="CV Meta")
        keep(m, True)
    return p

def add_bullet(text, bold_lead=None):
    p = doc.add_paragraph(style="CV Bullet")
    p.add_run("•  ")
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    keep(p)
    return p

doc.add_paragraph("LI HONG YI", style="Title")
doc.add_paragraph("Technical Artist  |  Creative Technologist  |  3D Artist", style="CV Subtitle")

add_section("EDUCATION")
add_role("Sichuan Fine Arts Institute", "2018-2022", "Chongqing, China")
p = doc.add_paragraph()
p.add_run("Bachelor of Fine Arts in Film and Television Production").bold = True
p.add_run("  |  GPA 3.22")
keep(p, True)
add_bullet("Core Courses: Film Theory, Audiovisual Language, Film and Television Production, Digital Media Production, 3D Visualization, Virtual Visual Effects", "Core Courses:")

add_section("WORK EXPERIENCE")
add_role("Applied R&D (Technical Artist)", "Aug 2023-Mar 2026", "Tencent Games - LIGHTSPEED STUDIOS, Omni Research", "Shenzhen Tengyu Interactive Technology Co., Ltd. (a Tencent subsidiary)")
doc.add_paragraph("KEY RESPONSIBILITIES", style="Heading 3")
for text in [
    "Contributed to the design and continuous iteration of batch-processing pipelines for 3D data in AI auto-rigging projects.",
    "Managed data production and quality control, including task planning, team training, quality review, issue tracking, and final delivery.",
    "Evaluated AI auto-rigging results, analyzed recurring issues across 3D models of varying complexity, and contributed to post-processing features and plugin interactions.",
    "Evaluated existing AI tools for capabilities, performance, and interaction workflows; created and maintained data-production guidelines, workflow documentation, and tool documentation.",
    "Provided 3D modeling, animation processing, rendering, and other visual-content support for multiple AI research projects.",
]: add_bullet(text)

doc.add_paragraph("SELECTED CONTRIBUTIONS", style="Heading 3")
add_bullet("Peacekeeper Elite (和平精英) - Developed and continuously iterated batch-processing pipelines for an AI auto-rigging project; evaluated internal AI tools and contributed to post-processing features, plugin interactions, and integration into practical art-production workflows.", "Peacekeeper Elite (和平精英)")
add_bullet("Aoxing Heatwave (奥星热浪) - Managed motion-capture data production for an AI-generated dance project, including requirement coordination, annotation-team management, quality control, data delivery, and cross-team collaboration.", "Aoxing Heatwave (奥星热浪)")

add_role("3D Artist", "Oct 2022-Apr 2023", "Chongqing Yujitu Technology Co., Ltd. (重庆禹迹图科技有限公司)")
add_bullet("Created production-ready 3D models for commercial projects, following project specifications and asset requirements.")

add_section("RESEARCH CONTRIBUTIONS")
add_role("Light-SQ", None, "Structure-Aware Shape Abstraction with Superquadrics for Generated Meshes", "SIGGRAPH Asia 2025  |  Co-author")
add_bullet("Contributed to the development of a batch-rendering script.")
add_bullet("Designed and produced the project cover and poster.")

add_role("TailorRig", None, "Controllable and Structured Auxiliary Skeleton Generation for Garment-Aware Rigging")
add_bullet("Contributed continuously to dataset development and data production.")
add_bullet("Designed and produced the project cover, poster, and related visual materials.")

add_section("TECHNICAL SKILLS")
for label, value in [
    ("Skills", "3D Modeling, Procedural Materials, Real-Time Rendering, Shader Development, Data Processing, Pipeline Development, Workflow Automation"),
    ("Software", "Unreal Engine, Blender, Substance 3D Designer, Substance 3D Painter, Marvelous Designer, ZBrush, MotionBuilder"),
    ("Programming Languages", "Python, HLSL"),
]:
    p = doc.add_paragraph()
    p.add_run(label + ": ").bold = True
    p.add_run(value)
    keep(p)

add_section("RESEARCH INTERESTS")
p = doc.add_paragraph("Human-AI Collaboration, Creative Tools, Computational Media, Authoring Systems, Computer Graphics, Procedural Materials, Real-Time Rendering")
keep(p)

for para in doc.paragraphs:
    para.paragraph_format.widow_control = True

doc.core_properties.title = "Li Hong Yi Curriculum Vitae"
doc.core_properties.subject = "Technical Artist and Creative Technologist"
doc.save(OUTPUT)
print(OUTPUT)
